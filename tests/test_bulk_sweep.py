"""Прогон необязательного набора data/bulk/ целиком: 100 документов, 10 категорий товара.

Пять выданных файлов проверяют умения, этот набор нужен, чтобы заметить разбор, заточенный
под конкретный файл. Тесты не требуют ни базы, ни модели эмбеддингов: гоняется только разбор.
Нет папки data/bulk/ — тесты пропускаются, набор на приёмку не влияет.

Полноту текста считает независимый разбор (pypdf, openpyxl), а не тот, что проверяется.
"""

import re
from collections import Counter, defaultdict
from functools import cache
from pathlib import Path

import openpyxl
import pytest
from app.parsers.models import ParseOutcome
from app.services.ingest import prepare_fragments
from pypdf import PdfReader

BULK = Path(__file__).resolve().parents[1] / "data" / "bulk"
SUFFIXES = {".pdf", ".docx", ".xlsx"}
FILES = sorted(p for p in BULK.iterdir() if p.suffix in SUFFIXES) if BULK.is_dir() else []
FULL_SET = len(FILES) == 100

PASSPORT_SECTIONS = {"ПАСПОРТ ИЗДЕЛИЯ", "Технические характеристики"}
OFFER_SECTIONS = {"Прайс-лист (опт)", "Условия", "Ваш менеджер"}
# Замер на 100 файлах: доля слов исходника в фрагментах не ниже 0.897 (теряется колонтитул).
MIN_RECALL = 0.85
# Потерянное слово должно быть системным (колонтитул), а не относиться к одному товару.
# Замер: самое редкое потерянное слово встречается в 15 файлах.
SYSTEMATIC_LOSS = 6
TOKEN = re.compile(r"[0-9a-zа-яё]+(?:[.,][0-9]+)?", re.IGNORECASE)
ARTICLE = re.compile(r"\b[A-Z]{2,6}-[A-Z0-9]+\b")

pytestmark = pytest.mark.skipif(not FILES, reason="необязательный набор data/bulk/ не найден")
needs_full_set = pytest.mark.skipif(not FULL_SET, reason="в data/bulk/ не все 100 файлов")


@pytest.fixture(scope="module")
def parsed() -> dict[Path, ParseOutcome]:
    return {path: prepare_fragments(path) for path in FILES}


def _kind(path: Path) -> str:
    return path.stem.split("_")[1]


def _category(path: Path) -> str:
    return path.stem.split("_")[0]


def _joined(outcome: ParseOutcome) -> str:
    return "\n".join(item.text for item in outcome.fragments)


@cache
def _raw_text(path: Path) -> str:
    if path.suffix == ".pdf":
        return " ".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
    return " ".join(
        str(cell)
        for sheet in openpyxl.load_workbook(path, data_only=True).worksheets
        for row in sheet.iter_rows(values_only=True)
        for cell in row
        if cell is not None
    )


def _tokens(text: str) -> set[str]:
    found = {token.lower().replace(",", ".") for token in TOKEN.findall(text)}
    return {token for token in found if len(token) >= 3 or any(ch.isdigit() for ch in token)}


def _product_rows(path: Path) -> int:
    """Строки товаров в xlsx: непустые строки каждой вкладки без заголовка."""
    total = 0
    for sheet in openpyxl.load_workbook(path, data_only=True).worksheets:
        rows = [r for r in sheet.iter_rows(values_only=True) if any(c is not None for c in r)]
        total += max(len(rows) - 1, 0)
    return total


def test_every_file_parses_with_page_and_section(parsed: dict[Path, ParseOutcome]) -> None:
    for path, outcome in parsed.items():
        assert outcome.reason is None, path.name
        assert outcome.fragments, path.name
        for item in outcome.fragments:
            assert item.text.strip(), path.name
            assert item.section, path.name
            assert item.page and item.page >= 1, path.name


@needs_full_set
def test_same_shape_in_every_category(parsed: dict[Path, ParseOutcome]) -> None:
    counts: dict[str, set[int]] = defaultdict(set)
    per_category: Counter[str] = Counter()
    for path, outcome in parsed.items():
        counts[_kind(path)].add(len(outcome.fragments))
        per_category[_category(path)] += len(outcome.fragments)
    assert len(counts["kp"]) == 1, counts["kp"]
    assert len(counts["spec"]) == 1, counts["spec"]
    assert len(set(per_category.values())) == 1, dict(per_category)


def test_spec_rows_stay_whole_with_article_and_brand(parsed: dict[Path, ParseOutcome]) -> None:
    specs = {path: outcome for path, outcome in parsed.items() if path.suffix == ".xlsx"}
    assert specs
    for path, outcome in specs.items():
        assert len(outcome.fragments) == _product_rows(path), path.name
        for item in outcome.fragments:
            assert item.article and item.brand, path.name
            assert item.text.startswith(f"{item.article}: "), path.name


def test_passports_and_offers_keep_sections_and_articles(
    parsed: dict[Path, ParseOutcome],
) -> None:
    for path, outcome in parsed.items():
        sections = {item.section for item in outcome.fragments}
        if _kind(path) == "passport":
            parts = path.stem.split("_")
            assert PASSPORT_SECTIONS <= sections, path.name
            assert f"{parts[2]}-{parts[3]}".upper() in _joined(outcome), path.name
        elif _kind(path) == "kp":
            assert OFFER_SECTIONS <= sections, path.name


def test_article_codes_of_the_source_are_kept(parsed: dict[Path, ParseOutcome]) -> None:
    for path, outcome in parsed.items():
        text = _joined(outcome)
        lost = {code for code in ARTICLE.findall(_raw_text(path)) if code not in text}
        assert not lost, (path.name, lost)


@needs_full_set
def test_source_text_survives_parsing(parsed: dict[Path, ParseOutcome]) -> None:
    lost_in: dict[str, list[str]] = defaultdict(list)
    for path, outcome in parsed.items():
        wanted = _tokens(_raw_text(path))
        missing = wanted - _tokens(_joined(outcome))
        assert 1 - len(missing) / len(wanted) >= MIN_RECALL, path.name
        for token in missing:
            lost_in[token].append(path.name)
    one_off = {t: names[:3] for t, names in lost_in.items() if len(names) < SYSTEMATIC_LOSS}
    assert not one_off, one_off
