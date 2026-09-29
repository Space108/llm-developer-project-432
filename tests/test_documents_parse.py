from pathlib import Path

from app.parsers.docx import parse_docx
from app.parsers.models import TextBlock
from app.parsers.pdf import parse_pdf
from app.parsers.xlsx import parse_xlsx
from app.services.chunk import chunk_text
from app.services.ingest import NO_TEXT_LAYER, prepare_fragments
from app.services.normalize import normalize_blocks
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter


def test_two_page_document_keeps_a_line_that_appears_once() -> None:
    blocks = normalize_blocks(
        [
            TextBlock(page=1, section="", text="уникальная строка\nКолонтитул"),
            TextBlock(page=2, section="", text="другая строка\nКолонтитул"),
        ]
    )
    texts = [block.text for block in blocks]
    assert "уникальная строка" in texts[0]
    assert "другая строка" in texts[1]
    assert all("Колонтитул" not in text for text in texts)


def test_hyphen_break_is_joined() -> None:
    blocks = normalize_blocks([TextBlock(page=1, section="", text="мощ-\nность 800 Вт")])
    assert blocks[0].text == "мощность 800 Вт"


def test_ligatures_and_spaces_collapse() -> None:
    blocks = normalize_blocks([TextBlock(page=1, section="", text="ﬁle   name")])
    assert blocks[0].text == "file name"


def test_text_chunks_overlap() -> None:
    assert chunk_text("abcdefghij", 4, 1) == ["abcd", "defg", "ghij"]


def test_xlsx_row_keeps_headers_and_is_not_split(tmp_path: Path) -> None:
    path = tmp_path / "spec.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Спецификация"
    sheet.append(["Артикул", "мощность", "чаша", "бренд"])
    sheet.append(["BLD-800", "800 Вт", "1.5 л", "Норд"])
    sheet.append(["BLD-900", "900 Вт", "2 л", "Норд"])
    book.save(path)
    book.close()

    rows = parse_xlsx(path)
    assert rows[0].text == "BLD-800: мощность 800 Вт; чаша 1.5 л; бренд Норд"
    assert rows[0].article == "BLD-800"
    assert rows[0].brand == "Норд"
    assert rows[0].page == 1
    assert rows[0].section == "Спецификация"
    outcome = prepare_fragments(path)
    assert outcome.reason is None
    assert len(outcome.fragments) == 2
    assert outcome.fragments[0].text == rows[0].text
    assert "BLD-900" not in outcome.fragments[0].text
    assert "BLD-800" not in outcome.fragments[1].text


def test_docx_heading_becomes_section(tmp_path: Path) -> None:
    path = tmp_path / "card.docx"
    document = Document()
    document.add_heading("Характеристики", level=1)
    document.add_paragraph("мощность 800 Вт")
    document.save(path)

    blocks = parse_docx(path)
    body = [block for block in blocks if "мощность" in block.text]
    assert body[0].section == "Характеристики"
    assert body[0].page == 1


def test_blank_pdf_is_rejected_with_a_reason(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as handle:
        writer.write(handle)
    outcome = prepare_fragments(path)
    assert outcome.fragments == []
    assert outcome.reason == NO_TEXT_LAYER


def test_broken_pdf_is_rejected_with_a_reason(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a pdf")
    outcome = prepare_fragments(path)
    assert outcome.fragments == []
    assert outcome.reason


def test_pdf_text_is_read_page_by_page(tmp_path: Path) -> None:
    path = tmp_path / "pages.pdf"
    path.write_bytes(_ascii_pdf(["Hello page one", "Hello page two"]))
    blocks = parse_pdf(path)
    assert [block.page for block in blocks] == [1, 2]
    assert "Hello page one" in blocks[0].text
    assert "Hello page two" in blocks[1].text


def test_long_table_row_stays_one_fragment(tmp_path: Path) -> None:
    path = tmp_path / "long.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["Артикул", "мощность"])
    sheet.append(["BLD-800", "800 Вт " + ("х" * 50)])
    book.save(path)
    book.close()
    outcome = prepare_fragments(path)
    assert outcome.reason is None
    assert len(outcome.fragments) == 1
    assert outcome.fragments[0].text.startswith("BLD-800: мощность 800 Вт")
    assert outcome.fragments[0].text.endswith("х" * 50)


def _ascii_pdf(pages: list[str]) -> bytes:
    """Минимальный pdf с текстом Helvetica. Кириллицу этот шрифт не несёт."""
    font = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    page_objects: list[bytes] = []
    content_objects: list[bytes] = []
    for line in pages:
        stream = f"BT /F1 12 Tf 72 100 Td ({line}) Tj ET".encode("ascii")
        header = b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n"
        content_objects.append(header + stream + b"\nendstream")
    kids: list[int] = []
    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    objects.append(b"")  # pages, filled after ids are known
    next_id = 3
    for content in content_objects:
        content_id = next_id
        page_id = next_id + 1
        font_id = next_id + 2
        next_id += 3
        kids.append(page_id)
        content_objects_ids = (content_id, page_id, font_id)
        page_objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] "
                f"/Contents {content_objects_ids[0]} 0 R "
                f"/Resources << /Font << /F1 {content_objects_ids[2]} 0 R >> >> >>"
            ).encode("ascii")
        )
        objects.append(content)
        objects.append(page_objects[-1])
        objects.append(font)
    kids_ref = " ".join(f"{kid} 0 R" for kid in kids)
    objects[1] = f"<< /Type /Pages /Kids [{kids_ref}] /Count {len(pages)} >>".encode("ascii")
    chunks = [b"%PDF-1.4\n"]
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(sum(len(chunk) for chunk in chunks))
        chunks.append(f"{index} 0 obj\n".encode("ascii") + body + b"\nendobj\n")
    xref_at = sum(len(chunk) for chunk in chunks)
    xref = [f"xref\n0 {len(objects) + 1}\n".encode("ascii"), b"0000000000 65535 f \n"]
    for offset in offsets[1:]:
        xref.append(f"{offset:010d} 00000 n \n".encode("ascii"))
    trailer = (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    ).encode("ascii")
    return b"".join(chunks + xref + [trailer])
