from pathlib import Path

from openpyxl import load_workbook

from app.parsers.models import TableRow

_ARTICLE_HEADERS = {"артикул", "article", "sku"}
_BRAND_HEADERS = {"бренд", "brand"}


def parse_xlsx(path: Path) -> list[TableRow]:
    """Первая строка листа — заголовки. Строка становится «артикул: параметр значение»."""
    book = load_workbook(path, read_only=True, data_only=True)
    rows: list[TableRow] = []
    try:
        for sheet_index, sheet in enumerate(book.worksheets, start=1):
            grid = [tuple(raw) for raw in sheet.iter_rows(values_only=True)]
            if not grid:
                continue
            headers = [_cell(value) for value in grid[0]]
            if not any(headers):
                continue
            article_at = _column(headers, _ARTICLE_HEADERS, default=0)
            brand_at = _column(headers, _BRAND_HEADERS, default=None)
            for raw in grid[1:]:
                values = [_cell(value) for value in raw]
                if not any(values):
                    continue
                article = values[article_at] if article_at < len(values) else ""
                brand = ""
                if brand_at is not None and brand_at < len(values):
                    brand = values[brand_at]
                parts: list[str] = []
                for index, header in enumerate(headers):
                    if index == article_at or not header:
                        continue
                    if index >= len(values) or not values[index]:
                        continue
                    parts.append(f"{header} {values[index]}")
                text = f"{article}: {'; '.join(parts)}" if parts else article
                rows.append(
                    TableRow(
                        page=sheet_index,
                        section=sheet.title or "",
                        article=article,
                        brand=brand,
                        text=text,
                    )
                )
    finally:
        book.close()
    return rows


def _cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _column(headers: list[str], names: set[str], default: int | None) -> int | None:
    for index, header in enumerate(headers):
        if header.casefold() in names:
            return index
    return default
